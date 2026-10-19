from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls

ROOT=Path(__file__).resolve().parent
AS=ROOT/'assets'; DOC=ROOT/'docs'
prs=Presentation(); prs.slide_width= Inches(13.333); prs.slide_height= Inches(7.5)
blank=prs.slide_layouts[6]
NAVY='#132446'; NAVY2='#1E3767'; CYAN='#2F9EAC'; TEAL='#1B7F84'; PURPLE='#7868D9'; GOLD='#F3C94F'; LIGHT='#F4F7FC'; MID='#D9E3F0'; TEXT='#20314F'; MUTED='#61718A'; RED='#C95652'; GREEN='#2F9476'; WHITE='#FFFFFF'; PALE='#EAF2FB'

def rgb(h):
    h=h.lstrip('#'); return RGBColor(int(h[:2],16),int(h[2:4],16),int(h[4:],16))

def bg(slide,color=LIGHT):
    # Use the actual slide background instead of a full-canvas rectangle to avoid any off-canvas geometry.
    fill=slide.background.fill
    fill.solid(); fill.fore_color.rgb=rgb(color)

def textbox(slide,x,y,w,h,text,size=16,color=TEXT,bold=False,font='Aptos',align=PP_ALIGN.LEFT,margin=.02, valign=MSO_ANCHOR.TOP, italic=False):
    sh=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=sh.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=Inches(margin); tf.margin_right=Inches(margin); tf.margin_top=Inches(margin); tf.margin_bottom=Inches(margin); tf.vertical_anchor=valign
    p=tf.paragraphs[0]; p.alignment=align; p.space_after=Pt(0)
    r=p.add_run(); r.text=text; r.font.name=font; r.font.size=Pt(size); r.font.bold=bold; r.font.italic=italic; r.font.color.rgb=rgb(color)
    return sh

def rect(slide,x,y,w,h,fill=WHITE,line=MID,rad=True):
    sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rad else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=rgb(fill); sh.line.color.rgb=rgb(line); sh.line.width=Pt(1)
    return sh

def line(slide,x1,y1,x2,y2,color=CYAN,width=2):
    sh=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    sh.line.color.rgb=rgb(color); sh.line.width=Pt(width); sh.line.end_arrowhead=True
    return sh

def header(slide,eyebrow,title,sub=None,dark=False):
    co='#7EE3E2' if dark else TEAL; titleco=WHITE if dark else NAVY; subco='#D6E1F2' if dark else MUTED
    textbox(slide,.55,.30,12.2,.25,eyebrow.upper(),9,co,True)
    textbox(slide,.55,.65,12.2,.55,title,27,titleco,True,font='Aptos Display')
    if sub: textbox(slide,.57,1.28,12,.38,sub,11,subco)

def footer(slide,num,dark=False):
    color='#9CB4D1' if dark else '#8391A6'
    textbox(slide,.55,7.13,10,.18,'DIGITECH  ·  GLUCOTWIN INDIA  ·  DIGITAL TWIN CHALLENGE 2026',7.5,color,True)
    textbox(slide,12.25,7.1,.5,.2,f'{num:02d}',8,color,True,align=PP_ALIGN.RIGHT)

def pill(slide,x,y,w,text,fill=PALE,color=NAVY):
    rect(slide,x,y,w,.32,fill,fill,True); textbox(slide,x+.07,y+.04,w-.14,.22,text,8.5,color,True,align=PP_ALIGN.CENTER)

def bullets(slide,x,y,w,items,size=15,color=TEXT,leading=.52):
    for i,item in enumerate(items):
        yy=y+i*leading
        # colored dot
        shape=slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(yy+.10), Inches(.09), Inches(.09))
        shape.fill.solid(); shape.fill.fore_color.rgb=rgb(CYAN); shape.line.fill.background()
        textbox(slide,x+.2,yy,w-.2,leading-.02,item,size,color)

# 1 cover
s=prs.slides.add_slide(blank); bg(s,NAVY)
# Decorative circles
for x,y,d,c in [(9.9,1.05,2.3,NAVY2),(10.8,1.9,1.45,CYAN),(9.5,3.0,1.0,PURPLE),(11.15,3.55,.55,GOLD)]:
    sh=s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d)); sh.fill.solid(); sh.fill.fore_color.rgb=rgb(c); sh.fill.transparency=12; sh.line.fill.background()
textbox(s,.7,.65,5,.35,'DIGITECH  /  DIGITAL TWIN CHALLENGE 2026',10,'#7EE3E2',True)
textbox(s,.7,1.45,8.6,1.25,'GlucoTwin India',42,WHITE,True,font='Aptos Display')
textbox(s,.74,2.78,8.2,.9,'An explainable digital twin for two-hour early warning of glucose excursions',22,'#D8E5F5',True)
rect(s,.75,4.1,6.65,1.05,NAVY2,'#34527E',True)
textbox(s,1.02,4.34,6.1,.55,'HISTORICAL EHR  +  LIVE-LIKE WEARABLE STREAM  →  RISK ESTIMATE',12,WHITE,True,align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE)
textbox(s,.75,5.65,7.1,.7,'Prototype uses synthetic data only. Research demonstration, not a medical device.',12,'#B9C8DC')
textbox(s,.75,6.85,6,.25,'Zeeshan Ali Sheikh  ·  Team DigiTech',9,'#91A6C3',True)
footer(s,1,True)

# 2 public-health motivation
s=prs.slides.add_slide(blank); bg(s); header(s,'01 · The problem','From isolated readings to a changing patient state','The challenge is to fuse historical context with a dynamic stream and forecast one concrete outcome.')
rect(s,.62,2.05,3.75,3.65,NAVY,NAVY)
textbox(s,.92,2.42,3.1,.3,'INDIA · DIABETES CONTEXT',10,'#7EE3E2',True)
textbox(s,.92,2.92,3.0,1.05,'90 million',29,WHITE,True,font='Aptos Display')
textbox(s,.92,4.02,3.0,.9,'estimated adults aged 20–79 living with diabetes in 2024',15,'#D6E2F2')
textbox(s,.92,5.18,3.05,.35,'IDF Diabetes Atlas, 11th ed. (2025)',8.5,'#9EB2CE')
rect(s,4.7,2.05,3.65,3.65,WHITE,MID)
textbox(s,5.0,2.4,3.0,.4,'THE CARE GAP',12,TEAL,True)
bullets(s,5.02,3.0,2.95,['A lab value is a snapshot.','Glucose and daily context change hour by hour.','A forecast needs time-aligned signals.'],13, TEXT,.72)
rect(s,8.7,2.05,3.95,3.65,WHITE,MID)
textbox(s,9.0,2.4,3.25,.4,'OUR NARROW SCOPE',12,PURPLE,True)
textbox(s,9.0,3.0,3.1,1.15,'Forecast a synthetic glucose threshold crossing within two hours.',19,NAVY,True)
textbox(s,9.0,4.48,3.15,.85,'A proof-of-concept, not a diagnosis or treatment engine.',12,MUTED)
textbox(s,.7,6.2,11.9,.38,'Source: IDF Diabetes Atlas (2025); WHO Diabetes Fact Sheet. Statistics are context, not model-training data.',9,MUTED)
footer(s,2)

# 3 target and workflow
s=prs.slides.add_slide(blank); bg(s); header(s,'02 · Problem statement','A forecast that is specific enough to test','The model only forecasts new threshold crossings when the current simulated glucose is below the threshold.')
steps=[('1','Twin state','Patient profile + latest stream'),('2','Forecast','P(crossing ≥180 mg/dL in 2h)'),('3','Review','Check signal and context'),('4','Human decision','Clinician remains in control')]
for i,(n,t,b) in enumerate(steps):
    x=.65+i*3.15
    rect(s,x,2.5,2.55,2.25,WHITE,MID)
    pill(s,x+.18,2.72,.42,n,fill=CYAN,color=WHITE)
    textbox(s,x+.18,3.22,2.18,.4,t,16,NAVY,True)
    textbox(s,x+.18,3.8,2.16,.7,b,11,MUTED)
    if i<3: line(s,x+2.58,3.62,x+3.0,3.62,CYAN,2)
rect(s,.7,5.25,11.95,.83,'#FFF6E5','#E7D49A')
textbox(s,.95,5.47,11.4,.38,'180 mg/dL is a prototype label chosen for this simulator — not a universal clinical alert threshold.',12,'#815A12',True,align=PP_ALIGN.CENTER)
footer(s,3)

# 4 two streams
s=prs.slides.add_slide(blank); bg(s); header(s,'03 · Feature fusion','Two streams make one patient-specific state','Static data explains baseline context; dynamic signals show what is changing now.')
rect(s,.7,2.05,5.45,3.8,WHITE,'#BBD4EF')
rect(s,7.15,2.05,5.45,3.8,WHITE,'#A9DDDE')
pill(s,1.0,2.35,1.55,'STATIC / EHR',fill='#DDEAFF',color=NAVY)
pill(s,7.45,2.35,1.65,'DYNAMIC STREAM',fill='#DDF6F3',color=TEAL)
textbox(s,1.0,2.95,4.6,.4,'Historical / patient context',18,NAVY,True)
bullets(s,1.05,3.55,4.45,['Age, BMI and diabetes duration','HbA1c and hypertension flag','Synthetic adherence factor'],13,TEXT,.59)
textbox(s,7.45,2.95,4.6,.4,'Hourly observations',18,NAVY,True)
bullets(s,7.5,3.55,4.45,['CGM value, trend and variability','Meal carbs + steps/activity','Heart rate, HRV and sleep','Time since meal + stress proxy'],13,TEXT,.51)
line(s,6.15,3.95,7.12,3.95,PURPLE,2)
rect(s,3.35,6.1,6.65,.5,NAVY,NAVY)
textbox(s,3.55,6.21,6.25,.22,'TIME ALIGNMENT  ·  FEATURE WINDOW AT TIME t  ·  NO FUTURE DATA INPUTS',9,WHITE,True,align=PP_ALIGN.CENTER)
footer(s,4)

# 5 architecture diagram
s=prs.slides.add_slide(blank); bg(s); header(s,'04 · System design','A transparent path from synthetic data to review','Architecture keeps inference-time features separate from retrospective outcome labels.')
s.shapes.add_picture(str(AS/'architecture.png'), Inches(.48), Inches(1.75), width=Inches(12.4), height=Inches(4.95))
footer(s,5)

# 6 synthetic data and labels
s=prs.slides.add_slide(blank); bg(s); header(s,'05 · Data and target','Reproducible simulation — no patient data required','The simulator generates patient-specific trajectories and labels each timestamp from the following two hours.')
rect(s,.7,2.0,5.6,3.75,WHITE,MID)
textbox(s,1.0,2.3,4.8,.4,'What is generated?',17,NAVY,True)
bullets(s,1.02,2.9,4.75,['240 synthetic patient profiles','14-day timeline; last 2 hours are label context','80,160 rows total across the full stream','Meals, activity, HR/HRV, sleep and CGM-like values'],12,TEXT,.52)
rect(s,6.8,2.0,5.8,3.75,NAVY,NAVY)
textbox(s,7.1,2.3,5.0,.4,'Target construction',17,WHITE,True)
textbox(s,7.1,2.95,4.9,.64,'At time t, use only current and past features.',13,'#D8E5F5')
textbox(s,7.1,3.75,4.9,.7,'Positive label: either t+1h or t+2h reaches ≥180 mg/dL.',15,'#7EE3E2',True)
textbox(s,7.1,4.65,4.9,.6,'Rows already ≥180 mg/dL are excluded from the early-warning target and surfaced separately.',11,'#D8E5F5')
textbox(s,.75,6.12,11.8,.35,'Future glucose is used only to create the retrospective label. It is not included in FEATURE_COLUMNS.',10,TEAL,True)
footer(s,6)

# 7 model & evaluation design
s=prs.slides.add_slide(blank); bg(s); header(s,'06 · AI/ML approach','Patient-level evaluation to reduce leakage','An Extra Trees baseline estimates probability; the operating threshold is tuned on validation data only.')
rect(s,.72,2.0,3.55,3.75,WHITE,MID)
textbox(s,1.0,2.3,3.0,.4,'MODEL',12,TEAL,True)
textbox(s,1.0,2.85,2.9,.7,'Extra Trees',21,NAVY,True)
bullets(s,1.0,3.7,2.9,['Nonlinear tabular baseline','Validation threshold tuned for F2','Global feature importance'],11,TEXT,.55)
rect(s,4.88,2.0,3.55,3.75,WHITE,MID)
textbox(s,5.15,2.3,3.0,.4,'SPLIT BY PATIENT ID',12,PURPLE,True)
for i,(name,n,col) in enumerate([('Train',153,CYAN),('Validation',39,PURPLE),('Test',48,GOLD)]):
    yy=3.02+i*.7
    rect(s,5.16,yy,2.8,.42,'#EEF2F8','#EEF2F8')
    rect(s,5.16,yy,2.8*n/153,.42,col,col)
    textbox(s,5.25,yy+.08,2.6,.22,f'{name} · {n} patients',10,NAVY,True)
rect(s,9.05,2.0,3.55,3.75,WHITE,MID)
textbox(s,9.35,2.3,2.85,.4,'THRESHOLD',12,'#A17C11',True)
textbox(s,9.35,2.85,2.9,.7,'F2 tuned',21,NAVY,True)
textbox(s,9.35,3.6,2.7,1.3,'Validation threshold chosen to emphasise recall for an early-warning demo; false alerts remain a trade-off.',12,TEXT)
textbox(s,.78,6.15,11.8,.35,'No patient can appear in more than one split. Metrics are estimates on a synthetic domain only.',10,MUTED)
footer(s,7)

# 8 metrics
s=prs.slides.add_slide(blank); bg(s); header(s,'07 · Measured results','Held-out synthetic test results','These numbers are computed by the included training script; they are not clinical performance claims.')
s.shapes.add_picture(str(AS/'model_metrics.png'), Inches(.58), Inches(1.78), width=Inches(7.2), height=Inches(4.25))
rect(s,8.0,2.05,4.6,3.8,NAVY,NAVY)
textbox(s,8.35,2.4,3.9,.35,'AT THE SELECTED THRESHOLD',10,'#7EE3E2',True)
textbox(s,8.35,3.0,3.75,.62,'68.5%',28,WHITE,True)
textbox(s,8.35,3.62,3.75,.4,'recall',13,'#D8E5F5')
textbox(s,8.35,4.17,3.75,.6,'25.9%',24,'#F3C94F',True)
textbox(s,8.35,4.78,3.75,.4,'precision',13,'#D8E5F5')
textbox(s,8.35,5.27,3.75,.36,'High recall comes with false alerts.',10,'#D8E5F5')
textbox(s,.75,6.35,11.6,.3,'Test set: 48 held-out patients, 15,184 eligible hourly rows. Positive event prevalence ≈ 4.5% overall.',9,MUTED)
footer(s,8)

# 9 dashboard screenshot preview
s=prs.slides.add_slide(blank); bg(s); header(s,'08 · Prototype','A clinician-facing, patient-specific dashboard','The UI connects the probability to a 24-hour trend, EHR context and a safe review workflow.')
s.shapes.add_picture(str(AS/'dashboard_preview.png'), Inches(.55), Inches(1.62), width=Inches(12.2), height=Inches(4.95))
textbox(s,.72,6.62,11.7,.22,'Visual preview rendered from the included synthetic data and trained model. The repository includes the working Streamlit application.',8.5,MUTED,italic=True)
footer(s,9)

# 10 what-if and explainability
s=prs.slides.add_slide(blank); bg(s); header(s,'09 · Explainability & interaction','Let reviewers explore sensitivity without claiming causality','The what-if panel changes model inputs; it does not simulate full human physiology or prove that an intervention causes a safer outcome.')
rect(s,.7,2.0,5.1,3.95,WHITE,MID)
textbox(s,1.0,2.3,4.4,.4,'WHAT-IF INPUTS',13,TEAL,True)
for i,(lab,fill,col) in enumerate([('Carbs last hour','#E4EEFC',NAVY),('Steps last hour','#DDF5F1',TEAL),('Sleep last night','#F2EFFE',PURPLE)]):
    yy=3.0+i*.78
    textbox(s,1.02,yy,1.65,.28,lab,11,TEXT,True)
    rect(s,2.85,yy-.02,2.2,.26,fill,fill,True)
    # slider track
    rect(s,2.97,yy+.075,1.85,.07,'#CED9E8','#CED9E8',True)
    circle=s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(3.7+i*.35), Inches(yy+.005), Inches(.22), Inches(.22)); circle.fill.solid(); circle.fill.fore_color.rgb=rgb(col); circle.line.fill.background()
textbox(s,1.02,5.5,4.35,.25,'Forecast updates when the scenario is applied.',10,MUTED)
rect(s,6.3,2.0,6.25,3.95,WHITE,MID)
s.shapes.add_picture(str(AS/'feature_importance.png'), Inches(6.6), Inches(2.35), width=Inches(5.65), height=Inches(3.05))
textbox(s,6.65,5.55,5.5,.25,'Global feature importance ≠ causality or a local explanation.',9,MUTED,True)
footer(s,10)

# 11 responsible use
s=prs.slides.add_slide(blank); bg(s); header(s,'10 · Responsible AI','Guardrails are part of the design','The prototype is deliberately positioned as a research demo and clinician review aid, not an automated decision maker.')
cols=[('WHAT IT DOES',['Joins two synthetic data streams','Estimates one narrowly scoped event','Shows metrics and data lineage'],CYAN),('WHAT IT DOES NOT DO',['Use real patient records or devices','Diagnose or recommend medication','Prove clinical efficacy or causality'],RED),('WHAT COMES NEXT',['External, governed datasets','Calibration + subgroup evaluation','Clinician workflow and safety review'],PURPLE)]
for i,(title,items,c) in enumerate(cols):
    x=.7+i*4.18
    rect(s,x,2.15,3.75,3.65,WHITE,MID)
    rect(s,x,2.15,3.75,.14,c,c,False)
    textbox(s,x+.25,2.55,3.2,.4,title,13,c,True)
    bullets(s,x+.26,3.25,3.05,items,12,TEXT,.67)
rect(s,.72,6.1,11.95,.55,'#FFF0EF','#F5CFCC')
textbox(s,.95,6.25,11.5,.2,'Synthetic data only · Not validated for real-world use · Never use to make diagnosis, treatment or emergency decisions.',10,'#9D322E',True,align=PP_ALIGN.CENTER)
footer(s,11)

# 12 close
s=prs.slides.add_slide(blank); bg(s,NAVY)
header(s,'11 · Conclusion','A small digital twin — end to end','GlucoTwin India demonstrates the complete challenge loop: patient state → hourly stream → forecast → clinician review.',dark=True)
items=[('01','FUSE','Historical + dynamic data'),('02','FORECAST','Two-hour event probability'),('03','EXPLAIN','Model results + limits'),('04','KEEP HUMAN','Clinician makes decisions')]
for i,(n,t,b) in enumerate(items):
    x=.7+i*3.15
    rect(s,x,2.45,2.65,1.78,NAVY2,'#34527E')
    textbox(s,x+.22,2.7,.45,.3,n,10,'#7EE3E2',True)
    textbox(s,x+.22,3.1,2.2,.36,t,15,WHITE,True)
    textbox(s,x+.22,3.56,2.17,.45,b,10,'#D6E2F2')
rect(s,.75,5.1,11.9,.75,'#244371','#416290')
textbox(s,1.0,5.33,11.4,.25,'Next milestone: evaluate on governed real-world data with clinical partners before any real deployment.',12,WHITE,True,align=PP_ALIGN.CENTER)
textbox(s,.75,6.2,11.7,.35,'Thank you · DigiTech · GlucoTwin India',18,'#7EE3E2',True)
footer(s,12,True)

# 13 references
s=prs.slides.add_slide(blank); bg(s); header(s,'Appendix · Sources','Context sources and reproducibility','Claims about public-health context are sourced; all model data and evaluation in this repo are synthetic.')
refs=[('International Diabetes Federation — Diabetes Atlas, 11th Edition (2025)','https://diabetesatlas.org/media/uploads/sites/3/2025/10/IDF_Diabetes_Atlas_11th_Edition_2025_WEB.pdf','India estimate: 90.0 million adults aged 20–79 living with diabetes in 2024.'),('World Health Organization — Diabetes Fact Sheet','https://www.who.int/news-room/fact-sheets/detail/diabetes','Overview of diabetes, complications, detection and management.'),('Digital Twin Challenge 2026 — Unstop / Happiest Health','https://unstop.com/hackathons/crp-digital-twin-challenge-2026-happiest-health-1757873','Challenge scope, two data streams and repository/video submission checklist.')]
for i,(title,url,desc) in enumerate(refs):
    yy=2.05+i*1.25
    rect(s,.72,yy,11.9,.98,WHITE,MID)
    textbox(s,1.0,yy+.15,11.15,.25,title,12,NAVY,True)
    textbox(s,1.0,yy+.48,11.2,.22,url,8.5,TEAL)
    textbox(s,1.0,yy+.71,11.2,.18,desc,9,MUTED)
textbox(s,.8,6.15,11.6,.3,'Implementation artefacts, generated metrics and limitations are available in the public repository README and docs.',9,MUTED)
footer(s,13)

out=DOC/'GlucoTwin_India_Presentation.pptx'
prs.save(out)
print(out)
